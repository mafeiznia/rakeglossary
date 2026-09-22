"""Export glossary entries to CSV, XLSX, and TBX."""

from __future__ import annotations

import csv
import json
import xml.etree.ElementTree as ET
from pathlib import Path
from xml.dom import minidom

from app.core.logging import get_logger
from app.models.glossary_entry import GlossaryEntry

log = get_logger("services.export")

# ---------------------------------------------------------------------------
# Shared column definitions
# ---------------------------------------------------------------------------

_CSV_HEADERS = [
    "English Term",
    "Persian Term",
    "Persian Alternatives",
    "Persian Transliteration",
    "POS",
    "Category",
    "Context",
    "Translator Note",
    "English Definition",
    "Persian Definition",
    "Source",
    "Score",
    "Frequency",
]

# CSV uses a compact 2-column layout; XLSX and TBX keep the full set.
_CSV_HEADERS_COMPACT = [
    "English Term",
    "Persian Term",
]


def _parse_alternatives(raw: str | None) -> list[str]:
    """Parse the JSON-encoded persian_alternatives string into a list."""
    if not raw:
        return []
    try:
        data = json.loads(raw)
        if isinstance(data, list):
            return [str(x).strip() for x in data if str(x).strip()]
    except (json.JSONDecodeError, TypeError):
        pass
    return []


def _row_values(e: GlossaryEntry) -> list:
    """Build the row of values in the same order as _CSV_HEADERS."""
    alternatives = _parse_alternatives(e.persian_alternatives)
    return [
        e.english_term,
        e.persian_term,
        " | ".join(alternatives),
        e.persian_transliteration or "",
        e.pos or "",
        e.category or "",
        e.context,
        e.translator_note or "",
        e.english_definition,
        e.persian_definition,
        e.source,
        e.score,
        e.frequency,
    ]


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


def export_csv(entries: list[GlossaryEntry], path: Path) -> Path:
    """Write entries to a UTF-8-SIG CSV file with 2 columns (term + translation).

    Intended for simple glossary exchange (e.g., importing into a CAT tool
    as a term pair list). For full metadata (context, POS, definitions, ...),
    use the XLSX or TBX export instead.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(_CSV_HEADERS_COMPACT)
        for e in entries:
            writer.writerow([e.english_term, e.persian_term])
    log.info(f"Exported {len(entries)} rows to CSV (compact 2-column): {path.name}")
    return path


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def export_xlsx(entries: list[GlossaryEntry], path: Path) -> Path:
    """Write entries to an XLSX workbook with RTL-friendly layout."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Glossary"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_align = Alignment(horizontal="center", vertical="center")

    for col, header in enumerate(_CSV_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Right-align Persian columns (2, 3, 4, 10 = Persian Term, Alternatives,
    # Transliteration, Persian Definition)
    rtl_columns = {2, 3, 4, 10}
    rtl_align = Alignment(horizontal="right", vertical="top", wrap_text=True)
    ltr_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

    for row_idx, e in enumerate(entries, start=2):
        values = _row_values(e)
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = rtl_align if col_idx in rtl_columns else ltr_align

    # Column widths (approximate, tunable)
    widths = [22, 22, 28, 20, 14, 22, 50, 40, 45, 45, 14, 8, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(path)
    log.info(f"Exported {len(entries)} rows to XLSX: {path.name}")
    return path


# ---------------------------------------------------------------------------
# TBX (TermBase eXchange)
# ---------------------------------------------------------------------------


def _tbx_lang_set(parent: ET.Element, lang: str) -> ET.Element:
    return ET.SubElement(parent, "langSet", {"xml:lang": lang})


def _tbx_tig(lang_set: ET.Element) -> ET.Element:
    return ET.SubElement(lang_set, "tig")


def _tbx_term(tig: ET.Element, text: str) -> None:
    if not text:
        return
    ET.SubElement(tig, "term").text = text


def _tbx_descrip(tig: ET.Element, dtype: str, text: str | None) -> None:
    if not text or not text.strip():
        return
    ET.SubElement(tig, "descrip", {"type": dtype}).text = text


def export_tbx(entries: list[GlossaryEntry], path: Path) -> Path:
    """Write entries to a TBX-Basic (TermBase eXchange) file."""
    path.parent.mkdir(parents=True, exist_ok=True)

    root = ET.Element(
        "tbx",
        {
            "style": "dct",
            "type": "TBX-Basic",
            "xml:lang": "en",
            "xmlns": "urn:iso:std:iso:30042:ed-2",
        },
    )
    header = ET.SubElement(root, "header")
    ET.SubElement(header, "fileDesc").text = "RakeGlossary export"
    body = ET.SubElement(root, "body")

    for i, e in enumerate(entries, start=1):
        te = ET.SubElement(body, "termEntry", {"id": f"t{i}"})

        # --- English side ---
        en_ls = _tbx_lang_set(te, "en")
        en_tig = _tbx_tig(en_ls)
        _tbx_term(en_tig, e.english_term)
        _tbx_descrip(en_tig, "partOfSpeech", e.pos)
        _tbx_descrip(en_tig, "subjectField", e.category)
        _tbx_descrip(en_tig, "definition", e.english_definition)
        _tbx_descrip(en_tig, "context", e.context)

        # --- Persian side ---
        alternatives = _parse_alternatives(e.persian_alternatives)
        if e.persian_term or alternatives or e.translator_note:
            fa_ls = _tbx_lang_set(te, "fa")
            fa_tig = _tbx_tig(fa_ls)
            _tbx_term(fa_tig, e.persian_term)
            for alt in alternatives:
                _tbx_term(fa_tig, alt)
            _tbx_descrip(fa_tig, "transliteration", e.persian_transliteration)
            _tbx_descrip(fa_tig, "definition", e.persian_definition)
            _tbx_descrip(fa_tig, "note", e.translator_note)

    # Pretty-print
    xml_bytes = ET.tostring(root, encoding="utf-8")
    pretty = minidom.parseString(xml_bytes).toprettyxml(indent="  ", encoding="utf-8")
    path.write_bytes(pretty)
    log.info(f"Exported {len(entries)} rows to TBX: {path.name}")
    return path


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

EXPORTERS = {
    "csv": (export_csv, ".csv"),
    "xlsx": (export_xlsx, ".xlsx"),
    "tbx": (export_tbx, ".tbx"),
}


def export(
    entries: list[GlossaryEntry],
    fmt: str,
    output_dir: Path,
    basename: str,
) -> Path:
    """Export entries in the given format. Returns the written path."""
    fmt = fmt.lower()
    if fmt not in EXPORTERS:
        raise ValueError(f"Unsupported format '{fmt}'. Use one of: {list(EXPORTERS)}")
    exporter, ext = EXPORTERS[fmt]
    path = output_dir / f"{basename}{ext}"
    return exporter(entries, path)
