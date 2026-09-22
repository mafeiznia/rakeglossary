"""DOCX extractor using python-docx, including tables."""

from __future__ import annotations

from pathlib import Path

import docx

from app.core.logging import get_logger
from app.pipeline.extractors.base import DocumentExtractor

log = get_logger("extractor.docx")


class DocxExtractor(DocumentExtractor):
    """Extract paragraphs and table cells from a DOCX file."""

    supported_extensions = (".docx",)

    def extract(self, file_path: Path) -> str:
        if not file_path.exists():
            log.error(f"File not found: {file_path}")
            return ""
        try:
            document = docx.Document(str(file_path))
        except Exception:
            log.exception(f"Failed to open DOCX '{file_path.name}'")
            return ""

        parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts)
