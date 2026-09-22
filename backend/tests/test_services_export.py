"""Tests for export_service (CSV, XLSX, TBX) with AI extras."""
from __future__ import annotations

import json
from pathlib import Path

from app.models.glossary_entry import GlossaryEntry
from app.services.export_service import export_csv, export_tbx, export_xlsx


def _sample_entries() -> list[GlossaryEntry]:
    return [
        GlossaryEntry(
            id=1,
            project_id="p1",
            english_term="Yumiko",
            persian_term="یومیکو",
            english_definition="",
            persian_definition="",
            source="LLM",
            score=0.5,
            frequency=3,
            context="Yumiko walked through Kagoshima.",
            pos="proper_noun",
            category="Character/Place Name",
            persian_alternatives=json.dumps(["یومیکو سان"], ensure_ascii=False),
            translator_note="نام شخصیت اصلی",
            persian_transliteration="Yumiko",
        ),
        GlossaryEntry(
            id=2,
            project_id="p1",
            english_term="DNA",
            persian_term="دی‌ان‌ای",
            english_definition="A molecule.",
            persian_definition="یک مولکول.",
            source="In-Text",
            score=0.9,
            frequency=5,
            context="DNA carries genetic information.",
            # AI extras are None (Offline/Hybrid mode)
        ),
    ]


def test_export_csv_creates_file(tmp_path: Path) -> None:
    out = export_csv(_sample_entries(), tmp_path / "g.csv")
    assert out.exists()
    content = out.read_text(encoding="utf-8-sig")
    # CSV is now a compact 2-column export
    lines = content.strip().splitlines()
    assert lines[0] == "English Term,Persian Term"
    assert len(lines) == 3  # header + 2 entries
    assert "Yumiko" in content
    assert "یومیکو" in content
    assert "DNA" in content
    assert "دی‌ان‌ای" in content
    # AI extras must NOT appear in the CSV
    assert "Persian Alternatives" not in content
    assert "proper_noun" not in content
    assert "نام شخصیت اصلی" not in content


def test_export_csv_is_two_columns(tmp_path: Path) -> None:
    """CSV export must contain exactly 2 columns (term + translation)."""
    out = export_csv(_sample_entries(), tmp_path / "g.csv")
    content = out.read_text(encoding="utf-8-sig")
    for line in content.strip().splitlines():
        assert line.count(",") == 1, f"Expected exactly 2 columns: {line!r}"


def test_export_xlsx_creates_file(tmp_path: Path) -> None:
    out = export_xlsx(_sample_entries(), tmp_path / "g.xlsx")
    assert out.exists()
    assert out.stat().st_size > 0

    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "English Term"
    assert ws.cell(row=1, column=3).value == "Persian Alternatives"
    assert ws.cell(row=1, column=5).value == "POS"
    assert ws.cell(row=2, column=1).value == "Yumiko"
    assert ws.cell(row=2, column=2).value == "یومیکو"
    assert ws.cell(row=2, column=5).value == "proper_noun"
    assert ws.cell(row=2, column=6).value == "Character/Place Name"
    assert ws.cell(row=2, column=8).value == "نام شخصیت اصلی"
    assert ws.cell(row=2, column=4).value == "Yumiko"


def test_export_xlsx_persian_columns_right_aligned(tmp_path: Path) -> None:
    out = export_xlsx(_sample_entries(), tmp_path / "g.xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb.active
    # Column 2 = Persian Term → right-aligned
    assert ws.cell(row=2, column=2).alignment.horizontal == "right"
    # Column 1 = English Term → left-aligned
    assert ws.cell(row=2, column=1).alignment.horizontal == "left"


def test_export_tbx_creates_file(tmp_path: Path) -> None:
    out = export_tbx(_sample_entries(), tmp_path / "g.tbx")
    assert out.exists()
    content = out.read_text(encoding="utf-8")
    assert "<tbx" in content
    assert 'xml:lang="en"' in content
    assert 'xml:lang="fa"' in content
    assert "Yumiko" in content
    assert "یومیکو" in content
    # AI extras
    assert "proper_noun" in content  # pos
    assert "Character/Place Name" in content  # category
    assert "نام شخصیت اصلی" in content  # translator_note
    # Alternative term
    assert "یومیکو سان" in content


def test_export_tbx_handles_missing_ai_extras(tmp_path: Path) -> None:
    """Entry 2 has no AI extras; TBX should still be valid."""
    out = export_tbx(_sample_entries(), tmp_path / "g.tbx")
    content = out.read_text(encoding="utf-8")
    assert "DNA" in content
    assert "دی‌ان‌ای" in content
    # Should not contain "None" as text
    assert ">None<" not in content


def test_export_csv_empty_entries(tmp_path: Path) -> None:
    out = export_csv([], tmp_path / "empty.csv")
    content = out.read_text(encoding="utf-8-sig")
    assert "English Term" in content
    assert len(content.strip().splitlines()) == 1  # header only


def test_export_xlsx_empty_entries(tmp_path: Path) -> None:
    out = export_xlsx([], tmp_path / "empty.xlsx")
    assert out.exists()
    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "English Term"
    # No data rows
    assert ws.cell(row=2, column=1).value is None


def test_export_tbx_invalid_alternatives_json(tmp_path: Path) -> None:
    """Broken JSON in persian_alternatives should not crash export."""
    entries = _sample_entries()
    entries[0].persian_alternatives = "not valid json {"
    out = export_tbx(entries, tmp_path / "g.tbx")
    assert out.exists()
    content = out.read_text(encoding="utf-8")
    assert "Yumiko" in content